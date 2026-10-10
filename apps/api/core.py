"""Core pipeline: chunk -> retrieve -> draft (cited) -> critic.
Designed and developed by NIKHIL CHARY SRIRAMOJU."""
import io, os, re, math, collections, datetime
from openpyxl import load_workbook

STOP = set("the a an of to and or in on for is are do does you your we our with be by as at it this that".split())
SENSITIVE = re.compile(r"breach|incident|certif|soc ?2 type|iso ?27001|legal|liabil|indemn|insurance", re.I)
tok = lambda t: [w for w in re.findall(r"[a-z0-9\-]+", t.lower()) if w not in STOP]


def chunk(text, size=600):
    out, buf = [], ""
    for p in (p.strip() for p in re.split(r"\n\s*\n", text) if p.strip()):
        if buf and len(buf) + len(p) > size:
            out.append(buf); buf = p
        else:
            buf = (buf + "\n" + p).strip()
    return out + ([buf] if buf else [])


class Index:
    def __init__(self): self.chunks = []; self.docs = {}

    def add(self, doc, text):
        m = re.search(r"reviewed:\s*(\d{4}-\d{2}-\d{2})", text, re.I)
        if m: self.docs[doc] = m.group(1)
        for i, c in enumerate(chunk(text)):
            self.chunks.append({"id": f"{doc}#{i+1}", "doc": doc, "text": c, "toks": tok(c)})

    def search(self, q, k=3):
        qt, n = set(tok(q)), len(self.chunks) + 1
        df = collections.Counter(t for c in self.chunks for t in set(c["toks"]))
        scored = []
        for c in self.chunks:
            tf = collections.Counter(c["toks"])
            s = sum((1 + math.log(tf[t])) * math.log(1 + n / df[t]) for t in qt if t in tf)
            if s > 0: scored.append((s, c))
        scored.sort(key=lambda x: -x[0])
        return [{"id": c["id"], "text": c["text"], "score": round(s, 2)} for s, c in scored[:k]]


def draft(q, hits):
    if not hits: return {"answer": "Insufficient evidence", "citations": []}
    if os.getenv("ANTHROPIC_API_KEY"):
        try:
            import anthropic
            ctx = "\n\n".join(f"[{h['id']}] {mask_pii(h['text'])}" for h in hits)
            r = anthropic.Anthropic().messages.create(
                model=os.getenv("ANTHROPIC_MODEL", "claude-sonnet-5-5"), max_tokens=400,
                system="Answer the security questionnaire question using ONLY the context. Cite chunk ids in "
                       "brackets like [doc#1]. If the context does not support an answer, reply exactly: Insufficient evidence",
                messages=[{"role": "user", "content": f"Context:\n{ctx}\n\nQuestion: {q}"}])
            a = r.content[0].text.strip()
            return {"answer": a, "citations": sorted(set(re.findall(r"\[([^\]]+#\d+)\]", a)))}
        except Exception:
            pass  # fall back to offline extractive mode
    qt = set(tok(q))
    sents = re.split(r"(?<=[.!?])\s+", hits[0]["text"])
    best = sorted(sents, key=lambda s: -len(qt & set(tok(s))))[:2]
    return {"answer": " ".join(best), "citations": [hits[0]["id"]]}


def critic(q, ans, hits):
    if ans["answer"].startswith("Insufficient") or not ans["citations"]:
        return {"verdict": "escalate", "confidence": 0.0, "reason": "no supporting evidence"}
    ids = {h["id"]: h for h in hits}
    if any(c not in ids for c in ans["citations"]):
        return {"verdict": "escalate", "confidence": 0.1, "reason": "citation not in retrieved context"}
    src = {t for c in ans["citations"] for t in tok(ids[c]["text"])}
    at = tok(ans["answer"])
    sup = round(sum(t in src for t in at) / max(len(at), 1), 2)
    if SENSITIVE.search(q):
        return {"verdict": "escalate", "confidence": sup, "reason": "sensitive topic: security owner must approve"}
    return {"verdict": "pass" if sup >= 0.6 else "escalate", "confidence": sup,
            "reason": "supported by sources" if sup >= 0.6 else "weak support"}


def mask_pii(t):
    """Mask emails, phone numbers and IPs before text is sent to an LLM."""
    t = re.sub(r"[\w.+-]+@[\w-]+\.[\w.]+", "[EMAIL]", t)
    t = re.sub(r"\b\d{1,3}(\.\d{1,3}){3}\b", "[IP]", t)
    return re.sub(r"(?<!\w)\+?\d[\d\s().-]{8,}\d", "[PHONE]", t)


def reuse(q, lib, thr=0.8):
    a = set(tok(q))
    for e in lib or []:
        b = set(tok(e["question"]))
        if a and b and len(a & b) / len(a | b) >= thr: return e
    return None


def answer_question(idx, q, lib=None):
    e = reuse(q, lib)
    if e:
        return {"question": q, "answer": e["answer"], "citations": e["citations"], "verdict": "pass",
                "confidence": 1.0, "reason": "reused approved answer", "status": "reused"}
    hits = idx.search(q)
    a = draft(q, hits)
    return {"question": q, **a, **critic(q, a, hits), "status": "drafted"}


def stale_docs(idx, days=365, today=None):
    today = today or datetime.date.today()
    return [{"doc": d, "reviewed": r, "age_days": (today - datetime.date.fromisoformat(r)).days}
            for d, r in idx.docs.items() if (today - datetime.date.fromisoformat(r)).days > days]


def gap_analysis(items):
    esc = [i for i in items if i["verdict"] != "pass"]
    return {"escalated": len(esc), "by_reason": dict(collections.Counter(i["reason"] for i in esc)),
            "questions": [i["question"] for i in esc]}


def parse_questions(data: bytes):
    wb = load_workbook(io.BytesIO(data))
    out = []
    for ws in wb.worksheets:
        col = None
        for row in ws.iter_rows(min_row=1, max_row=5):
            for c in row:
                if isinstance(c.value, str) and "question" in c.value.lower():
                    col = c.column; break
            if col: break
        for row in ws.iter_rows(min_row=1):
            c = row[(col or 1) - 1]
            if isinstance(c.value, str) and len(c.value.strip()) > 10 and "question" not in c.value.lower()[:9]:
                out.append({"sheet": ws.title, "row": c.row, "col": c.column, "text": c.value.strip()})
    return out


def export_xlsx(data: bytes, items):
    wb = load_workbook(io.BytesIO(data))
    for it in items:
        ws = wb[it["sheet"]]
        ws.cell(it["row"], it["col"] + 1, it["answer"])
        ws.cell(it["row"], it["col"] + 2, ", ".join(it["citations"]))
        ws.cell(it["row"], it["col"] + 3, it["status"])
    buf = io.BytesIO(); wb.save(buf); return buf.getvalue()
